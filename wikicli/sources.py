"""Read original sources into citable passages.

A passage is the unit that retrieval returns and that answers cite. Each one keeps
its source path and location (slide / page / section), so every citation can be
opened and checked. Supported formats: .pptx (slides + speaker notes), .pdf (pages,
parsed locally), .md / .txt (split by headings).
"""
from __future__ import annotations

import hashlib
import re
from dataclasses import asdict, dataclass
from pathlib import Path

from .config import VAULT

SUPPORTED = {".pptx", ".pdf", ".md", ".txt"}

# Slide chrome that repeats on every slide and only adds noise to retrieval.
BOILERPLATE = [
    re.compile(r"^POLECON 156\s*[·|-].*$", re.I),
    re.compile(r"^PART \d+\s*[:·].*$", re.I),
    re.compile(r"^\d{1,2}$"),
]


@dataclass
class Passage:
    id: str            # e.g. "week-2-section:s8" (stable; used by tests and citations)
    kind: str          # "source" (original evidence) or "wiki" (generated note)
    source_id: str     # e.g. "week-2-section"
    label: str         # human label, e.g. "Week 2 Section"
    path: str          # path relative to the vault, e.g. "raw/POLECON156_Week2_DiscussionSection.pptx"
    locator: str       # e.g. "slide 8" or "slide 5 (part 2)"
    number: int        # slide/page number (0 if not applicable)
    title: str
    text: str

    def to_dict(self) -> dict:
        return asdict(self)

    @property
    def citation(self) -> str:
        return f"{self.label}, {self.locator}"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def source_label(path: Path) -> str:
    """Readable label from a filename: POLECON156_Week2_DiscussionSection.pptx -> 'Week 2 Section'."""
    stem = path.stem
    week = re.search(r"week\s*_?(\d+)", stem, re.I)
    if week:
        kind = "Reflection" if re.search(r"reflection", stem, re.I) else \
               "Section" if re.search(r"discussion|section", stem, re.I) else \
               "Lecture" if re.search(r"lecture|history|economy", stem, re.I) else "Notes"
        return f"Week {week.group(1)} {kind}"
    return re.sub(r"[_\-]+", " ", stem).strip().title()


def _clean_lines(text: str) -> list[str]:
    lines = []
    for line in text.splitlines():
        line = re.sub(r"\s+", " ", line.replace("•", "•")).strip()
        if line and not any(p.match(line) for p in BOILERPLATE):
            lines.append(line)
    return lines


# ---- format readers: each returns [(number, title, text)] -----------------------

def _iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group: recurse so grouped text boxes are not lost
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def read_pptx(path: Path) -> list[tuple[int, str, str]]:
    from pptx import Presentation

    units = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        chunks = []
        for sh in _iter_shapes(slide.shapes):
            if sh.has_text_frame and sh.text_frame.text.strip():
                chunks.append(sh.text_frame.text)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    chunks.append(" | ".join(c.text.strip() for c in row.cells))
        lines = _clean_lines("\n".join(chunks))
        notes = ""
        if slide.has_notes_slide:
            notes = " ".join(_clean_lines(slide.notes_slide.notes_text_frame.text))
        title = lines[0] if lines else f"Slide {n}"
        body = "\n".join(lines)
        if notes:
            body += f"\nSpeaker notes: {notes}"
        if len(" ".join(lines).split()) + len(notes.split()) >= 8:   # skip image-only slides
            units.append((n, title, body))
    return units


def read_pdf(path: Path) -> list[tuple[int, str, str]]:
    from pypdf import PdfReader

    units = []
    for n, page in enumerate(PdfReader(path).pages, 1):
        lines = _clean_lines(page.extract_text() or "")
        if lines:
            units.append((n, lines[0][:80], "\n".join(lines)))
    return units


def read_text(path: Path) -> list[tuple[int, str, str]]:
    units, title, buf, n = [], path.stem, [], 0
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if line.startswith("#"):
            if "".join(buf).strip():
                n += 1
                units.append((n, title, "\n".join(buf).strip()))
            title, buf = line.lstrip("# ").strip(), []
        else:
            buf.append(line)
    if "".join(buf).strip():
        units.append((n + 1, title, "\n".join(buf).strip()))
    return units


READERS = {".pptx": read_pptx, ".pdf": read_pdf, ".md": read_text, ".txt": read_text}
UNIT_NAME = {".pptx": "slide", ".pdf": "page", ".md": "section", ".txt": "section"}


def split_long(text: str, max_words: int) -> list[str]:
    """Split on line boundaries so no part exceeds max_words (a single long line stays whole)."""
    parts, cur, count = [], [], 0
    for line in text.splitlines():
        w = len(line.split())
        if cur and count + w > max_words:
            parts.append("\n".join(cur))
            cur, count = [], 0
        cur.append(line)
        count += w
    if cur:
        parts.append("\n".join(cur))
    return parts


def load_source(path: Path, max_words: int = 220) -> list[Passage]:
    ext = path.suffix.lower()
    if ext not in READERS:
        raise ValueError(f"Unsupported file type: {path.name} (supported: {', '.join(sorted(SUPPORTED))})")
    label = source_label(path)
    sid = slugify(label)
    rel = path.resolve().relative_to(VAULT.resolve()).as_posix()
    unit = UNIT_NAME[ext]
    passages = []
    for n, title, text in READERS[ext](path):
        parts = split_long(text, max_words)
        for i, part in enumerate(parts, 1):
            if i > 1:   # repeat the slide title so a continuation still has context
                part = f"{title}\n{part}"
            loc = f"{unit} {n}" + (f" (part {i})" if len(parts) > 1 else "")
            pid = f"{sid}:s{n}" + (f".{i}" if len(parts) > 1 and i > 1 else "")
            passages.append(Passage(pid, "source", sid, label, rel, loc, n, title, part))
    return passages


def discover(folder: Path) -> list[Path]:
    if folder.is_file():
        return [folder]
    return sorted(p for p in folder.iterdir() if p.suffix.lower() in SUPPORTED and not p.name.startswith("."))
